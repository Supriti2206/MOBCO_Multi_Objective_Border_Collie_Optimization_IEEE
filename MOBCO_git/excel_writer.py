"""
excel_writer.py -- Writes the per-problem and master-summary Excel
workbooks that back the IEEE paper's tables.

write_problem_workbook() produces one multi-sheet .xlsx per problem with:
  Run_Summary, Statistics, Parameters, Objective_Info, Iteration_History,
  Pareto_Solutions, Archive, Reference_Front, Expected_vs_Reality

write_master_summary() produces one workbook aggregating every problem's
summary statistics into a single comparison table (Master_Summary), plus a
flat per-run dump (All_Runs) and the experiment configuration used
(Configuration).
"""

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path

HEADER_FONT = Font(name='Arial')
BODY_FONT = Font(name='Arial')


def _style(ws, float_format='0.000000'):
    """Apply consistent IEEE-paper-friendly formatting to one worksheet:
    Arial font, centered/formatted header row, 6-decimal float formatting
    in the body, auto-sized (capped) column widths, and a frozen header
    row for scrolling."""
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            if isinstance(cell.value, float):
                cell.number_format = float_format
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None),
                    default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(width + 2, 10), 42)
    ws.freeze_panes = 'A2'


def _obj_labels(problem):
    """Build per-objective column labels like 'f1_Cost', 'f2_Yield' using
    the problem's optional 'objective_names', falling back to
    'f{i}_Objective_{i}' when no names were given."""
    names = problem.get('objective_names') or []
    out = []
    for i in range(problem['n_obj']):
        base = names[i] if i < len(names) else f'Objective_{i + 1}'
        out.append(f'f{i + 1}_{base}')
    return out


def write_problem_workbook(result, out_dir):
    """
    Write one MOBCO_<ProblemName>.xlsx workbook summarizing a single
    problem's full experiment (result comes from run_experiment.run_problem).
    Sheets (in order): Run_Summary, Statistics, Parameters, Objective_Info,
    Iteration_History, Pareto_Solutions, Archive, Reference_Front,
    Expected_vs_Reality.
    """
    problem = result['problem']
    config = result['config']
    run_outputs = result['run_outputs']
    run_metrics = result['run_metrics']
    agg = result['aggregate']
    ctx = result['metric_context']

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = out_dir / f"MOBCO_{problem['name']}.xlsx"

    n_obj = problem['n_obj']
    obj_cols = _obj_labels(problem)

    with pd.ExcelWriter(filename, engine='openpyxl') as writer:

        # 1. Run_Summary
        rows = []
        for i, m in enumerate(run_metrics):
            rows.append({
                'Run': i + 1,
                'Hypervolume': m['hv'],
                'IGD': m['igd'],
                'GD': m['gd'],
                'Spacing': m['spacing'],
                'Spread': m['spread'],
                'Epsilon': m['epsilon'],
                'Runtime': m['runtime'],
            })
        df_runs = pd.DataFrame(rows, columns=['Run', 'Hypervolume', 'IGD', 'GD',
                                              'Spacing', 'Spread', 'Epsilon',
                                              'Runtime'])
        df_runs.to_excel(writer, sheet_name='Run_Summary', index=False)

        # 2. Statistics 
        stat_rows = []
        spec = [('Hypervolume', 'hv', True), ('IGD', 'igd', False),
                ('GD', 'gd', False), ('Spacing', 'spacing', False),
                ('Spread', 'spread', False), ('Epsilon', 'epsilon', False),
                ('Front_Size', 'size', True)]
        for label, key, larger_better in spec:
            stat_rows.append({
                'Metric': label,
                'Better': 'higher' if larger_better else 'lower',
                'Mean': agg[f'{key}_mean'], 'Std': agg[f'{key}_std'],
                'Median': agg[f'{key}_median'], 'Best': agg[f'{key}_best'],
                'Worst': agg[f'{key}_worst'],
            })
        stat_rows.append({'Metric': 'Runtime', 'Better': 'lower',
                          'Mean': agg['runtime_mean'], 'Std': agg['runtime_std'],
                          'Median': float(np.median([m['runtime'] for m in run_metrics])),
                          'Best': float(np.min([m['runtime'] for m in run_metrics])),
                          'Worst': float(np.max([m['runtime'] for m in run_metrics]))})
        pd.DataFrame(stat_rows).to_excel(writer, sheet_name='Statistics', index=False)

        # 3. Parameters
        params = [
            ('Algorithm', 'MOBCO (Multi-Objective Border Collie Optimization)'),
            ('Problem', problem['name']),
            ('Description', problem['description']),
            ('Decision Variables (dim)', problem['dim']),
            ('Number of Objectives', n_obj),
            ('Objective Directions', ', '.join(problem['objective_types'])),
            ('Variable Bounds', f"{problem['bounds'][0]} ... {problem['bounds'][-1]}"),
            ('Population Size (N)', config.n_population),
            ('Number of Dogs (leaders)', config.n_dogs),
            ('External Archive Size', config.archive_size),
            ('Max Iterations', config.max_iterations),
            ('Independent Runs', config.n_runs),
            ('Base Seed', config.base_seed),
            ('Mutation', f'Gaussian jitter (p={config.mutation_rate}, '
                         f'sigma=0.1 x variable range)'),
            ('Herding Mechanism', 'Dogs (leaders) accelerate towards a sibling '
                                  'leader; sheep are pulled towards the nearest '
                                  'dog with a stalking (far) / gathering (near) '
                                  'split; kinematics s = v*t + 0.5*a*t^2'),
            ('Stagnation Response', "'Eyeing' flag flips after 5 generations "
                                    "without archive growth, briefly reversing "
                                    "the sheep's acceleration sign to re-diversify"),
            ('Density Estimator', 'Crowding distance, mixed min/max aware'),
            ('Function Evaluations / run', run_outputs[0]['n_eval']),
            ('Reference Front', f"{result['reference_kind']} "
                                f"({len(result['reference_front'])} points)"),
            ('HV Reference Point', ', '.join(f'{v:.6g}' for v in ctx.ref_point)),
            ('HV Normalisation', 'volume fraction of the reference box, in [0,1]'),
        ]
        pd.DataFrame(params, columns=['Parameter', 'Value']).to_excel(
            writer, sheet_name='Parameters', index=False)

        # 4. Objective_Info 
        best_obs = []
        worst_obs = []
        F_best = run_outputs[result['best_run']]['archive_fitness']  # the single best (highest-HV) run's final archive
        for j, t in enumerate(problem['objective_types']):
            col = F_best[:, j]
            best_obs.append(col.min() if t == 'min' else col.max())
            worst_obs.append(col.max() if t == 'min' else col.min())
        pd.DataFrame({
            'Objective_Index': list(range(1, n_obj + 1)),
            'Objective_Name': obj_cols,
            'Direction': ['Minimize' if t == 'min' else 'Maximize'
                          for t in problem['objective_types']],
            'Best_In_Best_Run': best_obs,
            'Worst_In_Best_Run': worst_obs,
            'Reference_Best': ctx.best,
            'Reference_Worst': ctx.worst,
        }).to_excel(writer, sheet_name='Objective_Info', index=False)

        # 5. Iteration_History 
        h = result['history']
        if h is not None:
            it_df = pd.DataFrame({
                'Iteration': h['iteration'],
                'Mean_Hypervolume': h['hv'],
                'Mean_IGD': h['igd'],
                'Mean_GD': h['gd'],
                'Mean_Archive_Size': h['archive_size'],
                'Mean_Front_Size': h['front_size'],
            })
            for j, col in enumerate(obj_cols):
                it_df[f'{col}_Best'] = h['best'][:, j]
                it_df[f'{col}_Mean'] = h['mean'][:, j]
        else:
            it_df = pd.DataFrame({'Iteration': []})
        it_df.to_excel(writer, sheet_name='Iteration_History', index=False)

        # 6/7. Pareto_Solutions and Archive
        best_out = run_outputs[result['best_run']]

        def _sol_frame(X, F):
            if len(F) == 0:
                return pd.DataFrame({'Solution': []})
            dec = [f'x{i + 1}' for i in range(X.shape[1])]
            df = pd.DataFrame(X, columns=dec)
            for j, col in enumerate(obj_cols):
                df[col] = F[:, j]
            df.insert(0, 'Solution', range(1, len(df) + 1))
            df.insert(1, 'Best_Run', result['best_run'] + 1)
            return df

        _sol_frame(best_out['front_decisions'], best_out['front']).to_excel(
            writer, sheet_name='Pareto_Solutions', index=False)
        _sol_frame(best_out['archive_decisions'], best_out['archive_fitness']).to_excel(
            writer, sheet_name='Archive', index=False)

        #  8. Reference_Front (expected) 
        R = result['reference_front']
        df_ref = pd.DataFrame(R, columns=obj_cols)
        df_ref.insert(0, 'Point', range(1, len(df_ref) + 1))
        df_ref.insert(1, 'Source', result['reference_kind'])
        df_ref.to_excel(writer, sheet_name='Reference_Front', index=False)

        # 9. Expected_vs_Reality 
        # Compares the analytic/approximated reference front (R) against
        # the "reality" achieved -- the combined (union + non-dominated)
        # front across ALL runs -- per objective, direction-aware.
        comb = result['combined_front']
        cmp_rows = []
        for j, t in enumerate(problem['objective_types']):
            exp_best = R[:, j].min() if t == 'min' else R[:, j].max()
            exp_worst = R[:, j].max() if t == 'min' else R[:, j].min()
            got_best = comb[:, j].min() if t == 'min' else comb[:, j].max()
            got_worst = comb[:, j].max() if t == 'min' else comb[:, j].min()
            gap = (got_best - exp_best) if t == 'min' else (exp_best - got_best)
            cmp_rows.append({
                'Objective': obj_cols[j],
                'Direction': 'Minimize' if t == 'min' else 'Maximize',
                'Expected_Best': exp_best, 'Obtained_Best': got_best,
                'Expected_Worst': exp_worst, 'Obtained_Worst': got_worst,
                'Gap_To_Expected_Best': gap,
            })
        cmp_df = pd.DataFrame(cmp_rows)
        cmp_df.to_excel(writer, sheet_name='Expected_vs_Reality', index=False)

        for ws in writer.book.worksheets:
            _style(ws)

    return str(filename)


def write_master_summary(all_results, config, out_dir):
    """MOBCO_Master_Summary.xlsx — one row per problem (Master_Summary
    sheet) + the full run-level dump across all problems (All_Runs sheet)
    + the experiment's configuration (Configuration sheet). `all_results`
    is a dict of {problem_name: result} as produced by run_problem, for
    every problem in the sweep."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = out_dir / 'MOBCO_Master_Summary.xlsx'

    rows, long_rows = [], []
    for name, res in all_results.items():
        p, a = res['problem'], res['aggregate']
        rows.append({
            'Problem': p['name'],
            'Family': ('EvoPINN' if p['name'].startswith('EvoPINN')
                       else 'Mixed' if p['name'].startswith('Mixed')
                       else 'ZDT' if p['name'].startswith('ZDT') else 'DTLZ'),
            'Dim': p['dim'], 'N_Objectives': p['n_obj'],
            'Directions': ', '.join(p['objective_types']),
            'Reference_Front': res['reference_kind'],
            'HV_Mean': a['hv_mean'], 'HV_Std': a['hv_std'], 'HV_Best': a['hv_best'],
            'IGD_Mean': a['igd_mean'], 'IGD_Std': a['igd_std'],
            'GD_Mean': a['gd_mean'], 'Spacing_Mean': a['spacing_mean'],
            'Spread_Mean': a['spread_mean'], 'Epsilon_Mean': a['epsilon_mean'],
            'Front_Size_Mean': a['size_mean'],
            'Runtime_Mean_sec': a['runtime_mean'],
            'Runtime_Total_sec': a['runtime_total'],
            'Runs': config.n_runs, 'Iterations': config.max_iterations,
        })
        for i, m in enumerate(res['run_metrics']):
            long_rows.append({'Problem': p['name'], 'Run': i + 1,
                              'Hypervolume': m['hv'], 'IGD': m['igd'],
                              'GD': m['gd'], 'Spacing': m['spacing'],
                              'Spread': m['spread'], 'Epsilon': m['epsilon'],
                              'Runtime': m['runtime']})

    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name='Master_Summary', index=False)
        pd.DataFrame(long_rows).to_excel(writer, sheet_name='All_Runs', index=False)
        pd.DataFrame([{'Parameter': k, 'Value': str(v)}
                      for k, v in config.as_dict().items()]).to_excel(
            writer, sheet_name='Configuration', index=False)
        for ws in writer.book.worksheets:
            _style(ws)
    return str(filename)
